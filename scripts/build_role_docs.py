#!/usr/bin/env python3
"""Fill the bundled four-role worksheet and emit one DOCX per owner."""

from __future__ import annotations

import argparse
import copy
import json
import random
import re
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SKILL_DIR = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "Role work sheets.docx"

ROLE_SPECS = (
    ("discussion_leader", "Discussion Leader", "Sheet 1: Discussion Leader"),
    ("summarizer", "Summarizer", "Sheet 2: Summarizer"),
    ("word_master", "Word Master", "Sheet 3: Word Master"),
    ("structure_analyzer", "Structure Analyzer", "Sheet 4: Structure Analyzer"),
)

# Optional fallback mapping from a Chinese owner name to a renderer-safe romanization,
# e.g. {"张三": "Zhang San"}. A role's own `display_name` field takes precedence.
DISPLAY_NAMES: dict[str, str] = {}

PURPOSES = (
    ("describe", "Describe something"),
    ("explain_why", "Explain why"),
    ("explain_steps", "Explain how (steps)"),
    ("compare_contrast", "Compare and contrast"),
)

SUMMARY_STRUCTURES = (
    ("time_order", "Time/order"),
    ("space_order", "Space/order"),
    ("main_idea_details", "Main idea + details"),
    ("cause_effect", "Cause and effect"),
    ("compare_contrast", "Compare and contrast"),
)

ANALYZER_STRUCTURES = (
    ("story_plot", "Story/plot (beginning -> problem -> climax -> ending)"),
    ("sequence", "Sequence/chronological order (time order)"),
    ("cause_effect", "Cause and effect"),
    ("compare_contrast", "Compare and contrast"),
    ("problem_solution", "Problem and solution"),
    ("description_list", "Description/list (main idea + details)"),
)

WORD_CATEGORIES = (
    ("new", "Brand-new words"),
    ("striking", "Funny or striking words"),
    ("repeated", "Words that appear again and again"),
    ("specialist", "Topic/specialist words"),
    ("idioms", "Idioms or unusual phrases"),
)

LEADER_PARTICIPATION = (
    ("everyone_spoke", "Everyone spoke at least once"),
    ("stayed_on_topic", "We stayed on topic"),
    ("connected_to_experience", "We connected the text to our own ideas/experiences"),
    ("ended_with_big_question", "We ended with a prediction or big question"),
)

IMAGE_REFERENCE_RE = re.compile(
    r"(?:\bimage\s*\d+\b|图片\s*\d+|第\s*\d+\s*张图片)",
    re.IGNORECASE,
)


SCHEMA = {
    "title": "Article title",
    "date": "YYYY-MM-DD (optional; defaults to today)",
    "roles": {
        "discussion_leader": {
            "name": "Owner name",
            "display_name": "Romanized or renderer-safe owner name (optional)",
            "questions": ["Four open-ended questions"],
            "answerers": ["Full name of one group member per question; leave blank to auto-assign group members when meeting_completed is true"],
            "best_ideas": ["One likely best idea per question"],
            "passage_reference": "Printed page number and paragraph",
            "passage_reason": "Why the passage matters",
            "big_question": "Closing prediction or big question",
            "meeting_completed": "true only when the user wants a completed post-discussion worksheet",
            "participation": [
                "everyone_spoke", "stayed_on_topic", "connected_to_experience", "ended_with_big_question"
            ],
            "strongest_question": "My strongest question will be #... because...",
            "next_question": "Next time I would also ask...",
        },
        "summarizer": {
            "name": "Owner name",
            "display_name": "Romanized or renderer-safe owner name (optional)",
            "topic_sentence": "One sentence",
            "purposes": ["describe", "explain_why", "explain_steps", "compare_contrast"],
            "sections": [
                {"heading": "Part name", "range": "P1-P4", "main_idea": "...", "fact": "..."}
            ],
            "structures": ["main_idea_details", "cause_effect"],
            "full_summary": "Structured English summary",
            "short_summary": "At most 40 English words and contains both key terms",
            "key_terms": ["term 1", "term 2"],
        },
        "word_master": {
            "name": "Owner name",
            "display_name": "Romanized or renderer-safe owner name (optional)",
            "categories": ["new", "repeated", "specialist", "idioms"],
            "words": [
                {"word": "...", "location": "P1", "sentence": "Exact source sentence", "guess": "...", "meaning": "..."}
            ],
            "star": {
                "word": "...", "location": "...", "reason": "...",
                "memory_picture": "...", "new_sentence": "...", "question": "..."
            },
        },
        "structure_analyzer": {
            "name": "Owner name",
            "display_name": "Romanized or renderer-safe owner name (optional)",
            "structures": ["problem_solution", "sequence", "description_list"],
            "signal_words": "first, next, finally",
            "map_title": "Main idea",
            "map_parts": [
                {"label": "Part 1", "location": "P1-P4", "detail": "..."}
            ],
            "proof": "Our text is organized as...",
            "surprise": "A likely point to discuss is...",
            "effect": "How the structure affects meaning",
        },
    },
}


def fail(message: str) -> None:
    raise ValueError(message)


def required(mapping: dict, key: str, where: str) -> object:
    value = mapping.get(key)
    if value is None or value == "" or value == []:
        fail(f"Missing {where}.{key}")
    return value


def require_list(mapping: dict, key: str, where: str, *, exact: int | None = None,
                 minimum: int | None = None, maximum: int | None = None) -> list:
    value = required(mapping, key, where)
    if not isinstance(value, list):
        fail(f"{where}.{key} must be a list")
    if exact is not None and len(value) != exact:
        fail(f"{where}.{key} must contain exactly {exact} items")
    if minimum is not None and len(value) < minimum:
        fail(f"{where}.{key} must contain at least {minimum} items")
    if maximum is not None and len(value) > maximum:
        fail(f"{where}.{key} must contain at most {maximum} items")
    return value


def require_nonempty_items(values: list, where: str) -> None:
    for index, value in enumerate(values, 1):
        if value is None or not str(value).strip():
            fail(f"{where}[{index}] must not be empty")


def require_page_reference(value: object, where: str) -> None:
    text = str(value)
    if IMAGE_REFERENCE_RE.search(text):
        fail(f"{where} must use a printed page number, not an image reference")
    if not re.search(r"(?:\bpages?\b\s*:?[ ]*\d|\b\d+\s*,\s*(?:p|paragraph)\s*\d)", text, re.IGNORECASE):
        fail(f"{where} must include a printed page number")


def role_display_name(role: dict) -> str:
    name = str(role["name"])
    return str(role.get("display_name") or DISPLAY_NAMES.get(name, name))


def format_page_field(value: object) -> str:
    text = str(value).strip()
    text = re.sub(r"^pages?\s*:?[ ]*", "", text, flags=re.IGNORECASE)
    return f"Page: {text}"


def validate(data: dict) -> None:
    required(data, "title", "root")
    roles = required(data, "roles", "root")
    if not isinstance(roles, dict):
        fail("root.roles must be an object")

    for key, _, _ in ROLE_SPECS:
        role = required(roles, key, "roles")
        if not isinstance(role, dict):
            fail(f"roles.{key} must be an object")
        required(role, "name", f"roles.{key}")

    leader = roles["discussion_leader"]
    questions = require_list(leader, "questions", "roles.discussion_leader", exact=4)
    answerers = require_list(leader, "answerers", "roles.discussion_leader", exact=4)
    best_ideas = require_list(leader, "best_ideas", "roles.discussion_leader", exact=4)
    require_nonempty_items(questions, "roles.discussion_leader.questions")
    require_nonempty_items(best_ideas, "roles.discussion_leader.best_ideas")
    for field in ("passage_reference", "passage_reason", "big_question", "strongest_question", "next_question"):
        required(leader, field, "roles.discussion_leader")
    require_page_reference(leader["passage_reference"], "roles.discussion_leader.passage_reference")
    participation = leader.get("participation")
    if participation is not None:
        if not isinstance(participation, list):
            fail("roles.discussion_leader.participation must be a list")
        unknown = set(participation) - {key for key, _ in LEADER_PARTICIPATION}
        if unknown:
            fail(f"Unknown leader participation item(s): {', '.join(sorted(unknown))}")
        if participation and not leader.get("meeting_completed", False):
            fail("Leader participation items may only be checked when meeting_completed is true")
    # Blank answerers are auto-assigned from the group when the meeting is completed.

    summarizer = roles["summarizer"]
    for field in ("topic_sentence", "full_summary", "short_summary"):
        required(summarizer, field, "roles.summarizer")
    purposes = require_list(summarizer, "purposes", "roles.summarizer", minimum=1)
    unknown = set(purposes) - {key for key, _ in PURPOSES}
    if unknown:
        fail(f"Unknown summarizer purpose(s): {', '.join(sorted(unknown))}")
    sections = require_list(summarizer, "sections", "roles.summarizer", exact=4)
    for index, section in enumerate(sections, 1):
        for field in ("heading", "range", "main_idea", "fact"):
            required(section, field, f"roles.summarizer.sections[{index}]")
        require_page_reference(section["range"], f"roles.summarizer.sections[{index}].range")
    structures = require_list(summarizer, "structures", "roles.summarizer", minimum=1)
    unknown = set(structures) - {key for key, _ in SUMMARY_STRUCTURES}
    if unknown:
        fail(f"Unknown summarizer structure(s): {', '.join(sorted(unknown))}")
    terms = require_list(summarizer, "key_terms", "roles.summarizer", exact=2)
    short_words = re.findall(r"\b[\w'-]+\b", str(summarizer["short_summary"]))
    if len(short_words) > 40:
        fail(f"roles.summarizer.short_summary has {len(short_words)} words; maximum is 40")
    lowered = str(summarizer["short_summary"]).lower()
    for term in terms:
        if str(term).lower() not in lowered:
            fail(f"Short summary must include key term: {term}")

    word_master = roles["word_master"]
    words = require_list(word_master, "words", "roles.word_master", minimum=4, maximum=6)
    for index, word in enumerate(words, 1):
        for field in ("word", "location", "sentence", "guess", "meaning"):
            required(word, field, f"roles.word_master.words[{index}]")
        require_page_reference(word["location"], f"roles.word_master.words[{index}].location")
    categories = require_list(word_master, "categories", "roles.word_master", minimum=1)
    unknown = set(categories) - {key for key, _ in WORD_CATEGORIES}
    if unknown:
        fail(f"Unknown word category/categories: {', '.join(sorted(unknown))}")
    star = required(word_master, "star", "roles.word_master")
    for field in ("word", "location", "reason", "memory_picture", "new_sentence", "question"):
        required(star, field, "roles.word_master.star")
    require_page_reference(star["location"], "roles.word_master.star.location")

    analyzer = roles["structure_analyzer"]
    analyzer_structures = require_list(analyzer, "structures", "roles.structure_analyzer", minimum=1)
    unknown = set(analyzer_structures) - {key for key, _ in ANALYZER_STRUCTURES}
    if unknown:
        fail(f"Unknown analyzer structure(s): {', '.join(sorted(unknown))}")
    for field in ("signal_words", "map_title", "proof", "surprise", "effect"):
        required(analyzer, field, "roles.structure_analyzer")
    parts = require_list(analyzer, "map_parts", "roles.structure_analyzer", exact=4)
    for index, part in enumerate(parts, 1):
        for field in ("label", "location", "detail"):
            required(part, field, f"roles.structure_analyzer.map_parts[{index}]")
        require_page_reference(part["location"], f"roles.structure_analyzer.map_parts[{index}].location")


def has_cjk(text: str) -> bool:
    return any("\u3400" <= char <= "\u9fff" for char in text)


def apply_font(run, text: str, size: float | None = None, bold: bool | None = None) -> None:
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if has_cjk(text):
        # Use a broadly available Unicode font so owner names render in the
        # bundled LibreOffice QA renderer as well as in Word.
        run.font.name = "STSong"
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{key}"), "STSong")


def replace_paragraph(paragraph, text: str, size: float = 12, bold: bool | None = None) -> None:
    old = paragraph.runs[0] if paragraph.runs else None
    old_name = old.font.name if old else None
    old_size = old.font.size if old else None
    old_bold = old.bold if old else None
    paragraph.clear()
    run = paragraph.add_run(text)
    if old_name:
        run.font.name = old_name
    apply_font(run, text, size=(old_size.pt if old_size else size), bold=(old_bold if bold is None else bold))


def replace_heading(paragraph, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(13.5)
    run.font.color.rgb = RGBColor(0x1A, 0x20, 0x29)
    run.bold = True


def remove_paragraph(paragraph) -> None:
    parent = paragraph._element.getparent()
    if parent is not None:
        parent.remove(paragraph._element)


def compact_paragraph(paragraph, size: float = 10.5) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for run in paragraph.runs:
        run.font.size = Pt(size)


def set_table_column_widths(table, widths: tuple[float, ...]) -> None:
    table.autofit = False
    grid = table._tbl.tblGrid
    if grid is not None:
        for grid_col, width in zip(grid.gridCol_lst, widths):
            grid_col.set(qn("w:w"), str(round(Inches(width).twips)))
    for column, width in zip(table.columns, widths):
        for cell in column.cells:
            cell.width = Inches(width)


def set_cell(cell, text: str, size: float = 9.2, bold: bool = False) -> None:
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            apply_font(run, text, size=size, bold=bold)


def checkbox_line(options: tuple[tuple[str, str], ...], selected: list[str]) -> str:
    chosen = set(selected)
    return "  ".join(f"{'☑' if key in chosen else '☐'} {label}" for key, label in options)


def fill_header(table, name: str, date: str, title: str, role: str, group: str) -> None:
    set_cell(table.cell(0, 0), f"Name: {name}", size=9.2)
    set_cell(table.cell(0, 1), f"Date: {date}", size=9.5)
    set_cell(table.cell(1, 0), f"Book/Text: {title}", size=8.6)
    set_cell(table.cell(1, 1), f"Book/Text: {title}", size=8.6)
    set_cell(table.cell(2, 0), f"My Role: {role}", size=9.2)
    set_cell(table.cell(2, 1), f"Group: {group}", size=8.0)


def assign_answerers(answerers: list[str], roles: dict, title: str) -> list[str]:
    """Fill blank 'Who answered?' entries with the other group members' full names."""
    leader_name = role_display_name(roles["discussion_leader"])
    others = [
        role_display_name(role)
        for key, role in roles.items()
        if key != "discussion_leader" and role_display_name(role) != leader_name
    ]
    random.Random(title).shuffle(others)
    filled: list[str] = []
    cursor = 0
    for answerer in answerers:
        if str(answerer).strip():
            filled.append(str(answerer).strip())
        else:
            filled.append(others[cursor % len(others)])
            cursor += 1
    return filled


def replace_in_runs(cell, old: str, new: str) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace(old, new)
    if old in cell.text:
        set_cell(cell, cell.text.replace(old, new), size=10.5, bold=True)


def build_summarizer_chart(doc, sections: list[dict], anchor_paragraph, empty_paragraph) -> None:
    """Replace the template's tab-separated header line with a real Details Chart table."""
    from docx.table import Table

    # The Discussion Leader answer table has the needed shape (header + four rows, three columns).
    chart_element = copy.deepcopy(doc.tables[1]._tbl)
    anchor_paragraph._element.addprevious(chart_element)
    chart = Table(chart_element, doc._body)
    set_table_column_widths(chart, (1.5, 2.15, 2.15))
    headers = ("Section/Heading", "Main idea (What is this part about?)", "1 key fact or piece of data")
    for cell, text in zip(chart.rows[0].cells, headers):
        set_cell(cell, text, size=9.2, bold=True)
    for row, section in zip(chart.rows[1:], sections):
        values = (
            f"{section['heading']}\n{section['range']}",
            section["main_idea"],
            section["fact"],
        )
        for cell, value in zip(row.cells, values):
            set_cell(cell, str(value), size=8.2)
    remove_paragraph(anchor_paragraph)
    remove_paragraph(empty_paragraph)


def fill_document(doc, data: dict) -> None:
    from datetime import date as date_type

    title = str(data["title"])
    assignment_date = str(data.get("date") or date_type.today().isoformat())
    roles = data["roles"]
    group = " / ".join(role_display_name(roles[key]) for key, _, _ in ROLE_SPECS)

    header_tables = (0, 2, 3, 5)
    for table_index, (key, label, _) in zip(header_tables, ROLE_SPECS):
        fill_header(doc.tables[table_index], role_display_name(roles[key]), assignment_date, title, label, group)

    for paragraph_index in (2, 9, 13, 19, 24, 27, 33, 39, 42, 54, 62, 76, 85, 97):
        replace_heading(doc.paragraphs[paragraph_index], doc.paragraphs[paragraph_index].text)
    replace_heading(doc.paragraphs[0], "Sheet 1: Discussion Leader")
    replace_heading(doc.paragraphs[22], "Sheet 2: Summarizer")
    replace_heading(doc.paragraphs[52], "Sheet 3: Word Master")
    replace_heading(doc.paragraphs[74], "Sheet 4: Structure Analyzer")

    leader = roles["discussion_leader"]
    for index, question in zip(range(5, 9), leader["questions"]):
        replace_paragraph(doc.paragraphs[index], str(question))
    replace_paragraph(doc.paragraphs[11], format_page_field(leader["passage_reference"]))
    replace_paragraph(doc.paragraphs[12], f"Why this passage is worth discussing: {leader['passage_reason']}")
    meeting_completed = bool(leader.get("meeting_completed", False))
    selected_participation = leader.get("participation")
    if selected_participation is None:
        selected_participation = [key for key, _ in LEADER_PARTICIPATION] if meeting_completed else []
    selected_participation = set(selected_participation)
    for paragraph_index, (key, label) in zip(range(15, 19), LEADER_PARTICIPATION):
        mark = "☑" if key in selected_participation else "☐"
        if key in selected_participation:
            note = ""
        elif meeting_completed:
            note = "  (not observed)"
        else:
            note = "  (complete after the meeting)"
        suffix = f": {leader['big_question']}" if key == "ended_with_big_question" else ""
        replace_paragraph(doc.paragraphs[paragraph_index], f"{mark} {label}{suffix}{note}")
    replace_paragraph(doc.paragraphs[20], str(leader["strongest_question"]))
    replace_paragraph(doc.paragraphs[21], str(leader["next_question"]))
    leader_table = doc.tables[1]
    set_table_column_widths(leader_table, (0.8, 1.4, 3.6))
    answerers = [str(item) for item in leader["answerers"]]
    if meeting_completed:
        answerers = assign_answerers(answerers, roles, title)
    replace_in_runs(leader_table.rows[0].cells[1], "initials", "name")
    for row, answerer, idea in zip(leader_table.rows[1:], answerers, leader["best_ideas"]):
        set_cell(row.cells[1], answerer, size=8.4)
        set_cell(row.cells[2], str(idea), size=8.2)
    for paragraph_index in (*range(14, 19), 20, 21):
        compact_paragraph(doc.paragraphs[paragraph_index])

    summarizer = roles["summarizer"]
    replace_paragraph(doc.paragraphs[25], f"What is this text explaining? (One sentence) {summarizer['topic_sentence']}")
    replace_paragraph(doc.paragraphs[26], "The author's purpose: " + checkbox_line(PURPOSES, summarizer["purposes"]))
    chart_header_paragraph = doc.paragraphs[28]
    chart_empty_paragraph = doc.paragraphs[30]
    replace_paragraph(doc.paragraphs[35], checkbox_line(SUMMARY_STRUCTURES, summarizer["structures"]))
    replace_paragraph(doc.paragraphs[38], str(summarizer["full_summary"]))
    replace_paragraph(doc.paragraphs[41], str(summarizer["short_summary"]))
    replace_paragraph(doc.paragraphs[43], "2 key terms your summary must include: " + " and ".join(map(str, summarizer["key_terms"])))

    word_master = roles["word_master"]
    selected_categories = set(word_master["categories"])
    for paragraph_index, (key, label) in zip(range(56, 61), WORD_CATEGORIES):
        replace_paragraph(doc.paragraphs[paragraph_index], f"{'☑' if key in selected_categories else '☐'} {label}")
    star = word_master["star"]
    replace_paragraph(doc.paragraphs[64], f"Word: {star['word']}    {format_page_field(star['location'])}")
    replace_paragraph(doc.paragraphs[65], f"I chose it because: {star['reason']}")
    replace_paragraph(doc.paragraphs[66], f"Draw it or sketch a memory picture:\n{star['memory_picture']}")
    replace_paragraph(doc.paragraphs[67], f"Use it in a new sentence (not from the book): {star['new_sentence']}")
    replace_paragraph(doc.paragraphs[68], f"A quick question for the group using this word: {star['question']}")
    for paragraph_index in range(63, 69):
        compact_paragraph(doc.paragraphs[paragraph_index])
    word_rows = doc.tables[4].rows[1:]
    for index, row in enumerate(word_rows, 1):
        if index <= len(word_master["words"]):
            item = word_master["words"][index - 1]
            values = (str(index), f"{item['word']}\n{item['location']}", item["sentence"], item["guess"], item["meaning"])
        else:
            values = (str(index), "", "", "", "")
        for cell, value in zip(row.cells, values):
            set_cell(cell, str(value), size=7.8)
    for row in reversed(word_rows[len(word_master["words"]):]):
        doc.tables[4]._tbl.remove(row._tr)

    analyzer = roles["structure_analyzer"]
    selected_structures = set(analyzer["structures"])
    for paragraph_index, (key, label) in zip(range(78, 84), ANALYZER_STRUCTURES):
        replace_paragraph(doc.paragraphs[paragraph_index], f"{'☑' if key in selected_structures else '☐'} {label}")
    replace_paragraph(doc.paragraphs[84], f"Signal words that helped me decide: {analyzer['signal_words']}")
    replace_paragraph(doc.paragraphs[87], "Completed structure map for this text:")
    replace_paragraph(doc.paragraphs[88], f"Title/Main idea: {analyzer['map_title']}")
    replace_paragraph(doc.paragraphs[89], "↓")
    map_paragraphs = (90, 92, 94, 96)
    arrow_paragraphs = (91, 93, 95)
    for paragraph_index, part in zip(map_paragraphs, analyzer["map_parts"]):
        replace_paragraph(doc.paragraphs[paragraph_index], f"{part['label']} ({part['location']}): {part['detail']}", bold=False)
    for paragraph_index in arrow_paragraphs:
        replace_paragraph(doc.paragraphs[paragraph_index], "↓")
    replace_paragraph(doc.paragraphs[98], str(analyzer["proof"]))
    replace_paragraph(doc.paragraphs[99], str(analyzer["surprise"]))
    replace_paragraph(doc.paragraphs[100], f"How does the structure affect the meaning? {analyzer['effect']}")
    for row, part in zip(doc.tables[6].rows[1:], analyzer["map_parts"]):
        values = (part["label"], part["location"], part["detail"])
        for cell, value in zip(row.cells, values):
            set_cell(cell, str(value), size=8.2)

    doc.core_properties.title = f"{title} - Role Work Sheets"
    doc.core_properties.subject = "English reading-circle role assignment"
    doc.core_properties.author = group

    # The source template contains one empty numbered paragraph after the word list.
    remove_paragraph(doc.paragraphs[61])

    # Insert the Summarizer chart last so the fixed paragraph indexes above stay valid.
    build_summarizer_chart(doc, summarizer["sections"], chart_header_paragraph, chart_empty_paragraph)


def element_text(element) -> str:
    return "".join(element.itertext())


def trim_trailing_empty_paragraphs(body) -> None:
    for child in reversed(list(body)):
        if child.tag.endswith("}sectPr"):
            continue
        if child.tag.endswith("}p") and not element_text(child).strip():
            body.remove(child)
            continue
        break


def split_master(master_path: Path, output_dir: Path, roles: dict) -> list[Path]:
    probe = Document(master_path)
    children = list(probe.element.body)
    heading_indices: dict[str, int] = {}
    for key, _, marker in ROLE_SPECS:
        matches = [index for index, child in enumerate(children) if marker in element_text(child)]
        if len(matches) != 1:
            fail(f"Could not uniquely locate template heading: {marker}")
        heading_indices[key] = matches[0]

    starts: dict[str, int] = {}
    for key, _, _ in ROLE_SPECS:
        index = heading_indices[key] - 1
        while index >= 0 and not children[index].tag.endswith("}tbl"):
            index -= 1
        if index < 0:
            fail(f"Could not locate header table before {key}")
        starts[key] = index

    output_paths: list[Path] = []
    for role_index, (key, label, _) in enumerate(ROLE_SPECS, 1):
        doc = Document(master_path)
        body = doc.element.body
        current_children = list(body)
        start = starts[key]
        if role_index < len(ROLE_SPECS):
            next_key = ROLE_SPECS[role_index][0]
            end = starts[next_key]
        else:
            end = len(current_children)

        for index, child in reversed(list(enumerate(current_children))):
            if child.tag.endswith("}sectPr"):
                continue
            if index < start or index >= end:
                body.remove(child)

        trim_trailing_empty_paragraphs(body)

        owner = safe_filename(str(roles[key]["name"]))
        role_name = label.replace(" ", "_")
        output_path = output_dir / f"{role_index}_{owner}_{role_name}.docx"
        doc.core_properties.title = f"{roles[key]['name']} - {label}"
        doc.save(output_path)
        output_paths.append(output_path)

    return output_paths


def safe_filename(value: str) -> str:
    cleaned = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", value).strip(" .")
    return cleaned or "owner"


def build(content_path: Path, output_dir: Path, template_path: Path) -> list[Path]:
    data = json.loads(content_path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        fail("Content JSON root must be an object")
    validate(data)
    if not template_path.is_file():
        fail(f"Template not found: {template_path}")
    output_dir.mkdir(parents=True, exist_ok=True)

    doc = Document(template_path)
    fill_document(doc, data)
    with tempfile.TemporaryDirectory(prefix="english_role_sheets_") as temp_dir:
        master_path = Path(temp_dir) / "all_roles.docx"
        doc.save(master_path)
        return split_master(master_path, output_dir, data["roles"])


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--content", type=Path, help="UTF-8 JSON containing grounded role content")
    parser.add_argument("--output-dir", type=Path, help="Directory for the four DOCX files")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE, help="Override the bundled DOCX template")
    parser.add_argument("--print-schema", action="store_true", help="Print the JSON content schema and exit")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.print_schema:
        print(json.dumps(SCHEMA, ensure_ascii=False, indent=2))
        return 0
    if args.content is None or args.output_dir is None:
        print("error: --content and --output-dir are required unless --print-schema is used", file=sys.stderr)
        return 2
    try:
        outputs = build(args.content.resolve(), args.output_dir.resolve(), args.template.resolve())
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1
    for output in outputs:
        print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
